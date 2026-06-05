#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
结果回填 —— 在本地运行（需要 高风险事件样本.xlsx 和 openpyxl）。

把服务器跑出的 results_*.jsonl 按 hash 映射回三张表的【所有】原始行
（包括同一视频被多行引用的情况），输出一张对比 CSV，便于和「标注结果」核对。

用法：
  python join_results.py \
      --xlsx   ../../高风险事件样本.xlsx \
      --results results_5.1-5.20.jsonl \
      --out     results_joined_5.1-5.20.csv
"""
import argparse
import csv
import json

import openpyxl

SHEETS = ["5.1~5.6", "5.7~5.13", "5.14~5.20"]


def load_results(path):
    res = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            r = json.loads(line)
            res[r["hash"]] = r
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--results", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    results = load_results(args.results)
    print(f"加载推理结果 {len(results)} 条")

    wb = openpyxl.load_workbook(args.xlsx, read_only=True)
    out = open(args.out, "w", newline="", encoding="utf-8")
    w = csv.writer(out)
    w.writerow(["sheet", "报警ID", "风险开始时间", "DMS模型结果", "标注结果",
                "hash", "pred_class", "pred_score", "stage", "infer_error"])

    n = matched = 0
    for sn in SHEETS:
        ws = wb[sn]
        header = list(next(ws.iter_rows(values_only=True)))
        iu = header.index("唤醒视频链接")
        ia = header.index("报警ID")
        it = header.index("风险开始时间")
        idms = header.index("DMS模型结果")
        ilab = header.index("标注结果（对应链接）")
        for row in ws.iter_rows(min_row=2, values_only=True):
            u = row[iu]
            if not u or not str(u).strip():
                continue
            n += 1
            h = str(u).rsplit("/", 1)[-1].replace(".mp4", "").split("_")[-1]
            r = results.get(h)
            if r:
                matched += 1
                w.writerow([sn, row[ia], row[it], row[idms], row[ilab], h,
                            r.get("class"), r.get("score"), r.get("stage"), r.get("error", "")])
            else:
                w.writerow([sn, row[ia], row[it], row[idms], row[ilab], h,
                            "", "", "", "未推理"])
    out.close()
    print(f"原始有链接行 {n}，已匹配模型结果 {matched}，写出 {args.out}")


if __name__ == "__main__":
    main()
